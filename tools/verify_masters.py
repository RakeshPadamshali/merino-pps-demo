"""Master Data, the full Excel round trip: download a template, edit it (an update, a new row, a broken row), upload it,
see the checks and the preview, apply it, and confirm the planning engine re-planned with the new master.
Usage:  python tools/verify_masters.py
"""
import os
import subprocess
import sys
import tempfile
import time

import openpyxl
from playwright.sync_api import sync_playwright

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PORT = 8422
BASE = "http://localhost:%d/" % PORT
TMP = tempfile.mkdtemp(prefix="mer-masters-")
fails = []


def check(cond, name, detail=""):
    print(("  ok   " if cond else "  FAIL ") + name + (("  — " + str(detail)) if detail and not cond else ""))
    if not cond:
        fails.append(name)


def ready(pg, timeout=90000):
    pg.wait_for_function("!document.getElementById('mer-busy') && !!window.MERX", timeout=timeout)
    pg.wait_for_timeout(300)


def open_master(pg, key, clear=False):
    pg.goto(BASE + "masters.html#" + key, wait_until="load")
    if clear:
        pg.evaluate("localStorage.removeItem('mer-demo-state')")
        pg.reload(wait_until="load")
    ready(pg)
    pg.click("[data-m='%s']" % key)
    pg.wait_for_timeout(300)


def write_xlsx(path, sheet, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def main():
    srv = subprocess.Popen([sys.executable, "-m", "http.server", str(PORT)], cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        with sync_playwright() as p:
            b = p.chromium.launch(channel="chrome", headless=True)
            ctx = b.new_context(viewport={"width": 1600, "height": 1000}, accept_downloads=True)
            pg = ctx.new_page()
            errs = []
            pg.on("pageerror", lambda e: errs.append(str(e)))
            pg.on("console", lambda m: errs.append("console: " + m.text) if m.type == "error" and "favicon" not in m.text else None)
            for i in range(25):
                try:
                    pg.goto(BASE + "masters.html", wait_until="load")
                    break
                except Exception:
                    time.sleep(0.4)

            # ---------- 1. template download ----------
            open_master(pg, "mouldPolicy", clear=True)
            with pg.expect_download() as dl:
                pg.click("#dl")
            d = dl.value
            tpl = os.path.join(TMP, d.suggested_filename)
            d.save_as(tpl)
            check(d.suggested_filename == "MER_mouldPolicy_template.xlsx", "template file is named after the master", d.suggested_filename)
            wb = openpyxl.load_workbook(tpl)
            check("mouldPolicy" in wb.sheetnames and "README" in wb.sheetnames, "template has the master sheet and a README", wb.sheetnames)
            ws = wb["mouldPolicy"]
            head = [c.value for c in ws[1]]
            check("id" in head and "value" in head, "template header carries the master's columns", head)
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            check(len(rows) == 5, "template holds the current rows", len(rows))
            readme = {r[0]: r[1] for r in wb["README"].iter_rows(values_only=True) if r and r[0]}
            check("yes" in str(readme.get("Drives the plan", "")), "README says the master drives the plan", readme.get("Drives the plan"))

            # ---------- 2. upload an edited file: one update, one new row, two broken rows ----------
            before = pg.evaluate("(() => { const m = MERX.model(); return { refurbMin: m.refurbMin, life: m.life, healthA: m.health.A }; })()")
            edit = os.path.join(TMP, "mouldPolicy_edit.xlsx")
            write_xlsx(edit, "mouldPolicy", ["id", "value", "unit", "meaning", "source"], [
                ["REFURB", 24, "hours", "Plate re-polish / re-chrome turnaround (express line)", "GENERATED"],   # update
                ["HEALTH-D", 2, "% life left", "Trial class for internal samples", "GENERATED"],                 # new
                ["BROKEN", "", "cycles", "no value at all", "GENERATED"],                                        # missing required value
                ["LIFE", "not-a-number", "cycles", "life limit", "GENERATED"],                                   # not numeric
            ])
            pg.set_input_files("#file", edit)
            pg.wait_for_timeout(700)
            check(pg.locator("#prevcard").is_visible(), "upload opens the preview")
            meta = pg.inner_text("#pmeta")
            check("1 new" in meta and "1 updates" in meta and "2 errors" in meta, "preview counts new / updated / rejected rows", meta)
            pv = pg.inner_text("#prev").lower()
            check("missing value" in pv, "a row without a required value is rejected with the reason", pv[:160])
            check("must be a number" in pv, "a non-numeric value is rejected with the reason")
            check("48 → 24" in pg.inner_text("#prev"), "the preview shows the old → new value of the updated row")
            check(pg.evaluate("(() => { const m = MERX.model(); return m.refurbMin; })()") == before["refurbMin"], "nothing is applied before the preview is accepted")

            # ---------- 3. apply → the engine re-plans with the new master ----------
            k0 = pg.evaluate("(() => { const st = MERX.settings(); const r = MERX.cachedPlan(st) || MERX.computePlan(st); return { otif: r.kpi.otif, cycles: r.kpi.cycles, sheets: r.kpi.sheets }; })()")
            pg.click("#apply")
            pg.wait_for_timeout(500)
            ready(pg)
            pg.wait_for_timeout(500)
            after = pg.evaluate("(() => { const m = MERX.model(); return { refurbMin: m.refurbMin, health: Object.keys(m.health).sort().join(',') }; })()")
            check(after["refurbMin"] == 24 * 60, "the applied row reaches the planning model (refurbishment 48 h → 24 h)", after)
            check("D" in after["health"], "the new health class row is in the model too", after["health"])
            k1 = pg.evaluate("(() => { const st = MERX.settings(); const r = MERX.cachedPlan(st) || MERX.computePlan(st); return { otif: r.kpi.otif, cycles: r.kpi.cycles, sheets: r.kpi.sheets }; })()")
            check(k1 != k0, "the press-shop plan was re-run with the new master", "%s → %s" % (k0, k1))
            audit = pg.inner_text("#audit")
            check("Mould policy" in audit and "ppsadmin" in audit and "48 → 24" in audit, "the audit row records who changed what", audit[:200])
            check("2 rows from Excel" in pg.inner_text("#mmeta"), "the master shows how many rows came from Excel", pg.inner_text("#mmeta"))

            # the change is visible on another page that reads the same model
            pg.goto(BASE + "moulds.html", wait_until="load")
            ready(pg)
            txt = pg.inner_text("#kpis")
            check("24 h" in txt, "the Mould Board shows the new refurbishment turnaround", txt[:200])

            # ---------- 4. a broken reference is refused ----------
            open_master(pg, "stockNorms")
            bad = os.path.join(TMP, "stockNorms_bad.xlsx")
            write_xlsx(bad, "stockNorms", ["id", "skuId", "normSheets", "usePerDay", "onHandAtBase"], [
                ["NRM-GHOST", "D9999-XX-1-84-GP", 300, 20, 100],
            ])
            pg.set_input_files("#file", bad)
            pg.wait_for_timeout(700)
            pv = pg.inner_text("#prev")
            check("not in" in pv and "0 new" in pg.inner_text("#pmeta"), "a row pointing at a SKU that does not exist is rejected", pv[:200])
            check(pg.locator("#apply").is_disabled(), "nothing can be applied when every row failed")
            pg.click("#cancel")

            # ---------- 5. reset ----------
            pg.goto(BASE + "masters.html", wait_until="load")
            ready(pg)
            pg.click("#mer-reset")
            pg.wait_for_timeout(500)
            ready(pg)
            st = pg.evaluate("localStorage.getItem('mer-demo-state')")
            check(not st, "Reset demo clears the uploaded masters", st)
            check(pg.evaluate("MERX.model().refurbMin") == before["refurbMin"], "the model is back to the shipped masters")
            check(not errs, "no JavaScript errors during the round trip", errs[:3])
            b.close()
    finally:
        srv.kill()
    print("\nMASTERS", "PASS" if not fails else "FAIL (%d)" % len(fails))
    sys.exit(0 if not fails else 1)


main()
